import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .config import default_data_root

from .types import (
    ArmyType,
    ArmyUnit,
    EquipmentItem,
    FactionPower,
    LandBlock,
    Minister,
    NavyUnit,
    NumericData,
    PartyGroup,
    ShipType,
    StrataGroup,
    TaxRule,
    TreasurySnapshot,
)


def _norm_text(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _parse_num(v, default=0.0) -> float:
    s = _norm_text(v)
    if not s or s == "-":
        return float(default)

    negative = False
    if s.startswith("-"):
        negative = True
        s = s[1:].strip()

    multiplier = 1.0
    if "亿" in s:
        multiplier *= 100000000.0
        s = s.replace("亿", "")
    if "万" in s:
        multiplier *= 10000.0
        s = s.replace("万", "")
    if "石" in s:
        s = s.replace("石", "")
    if "两" in s:
        s = s.replace("两", "")
    if "%" in s:
        s = s.replace("%", "")
        try:
            val = float(s) / 100.0
            return -val if negative else val
        except ValueError:
            return float(default)

    s = s.replace(",", "").replace(" ", "")
    try:
        val = float(s) * multiplier
        return -val if negative else val
    except ValueError:
        return float(default)


def _read_table(path: str) -> List[Dict[str, str]]:
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    rows = []
    for row in ws.iter_rows(min_row=1, max_row=200):
        rows.append([cell.value for cell in row])

    if not rows:
        return []

    header = [_norm_text(h) for h in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(cell is not None and _norm_text(cell) for cell in row):
            continue
        item = {}
        for idx, key in enumerate(header):
            if not key:
                continue
            item[key] = _norm_text(row[idx] if idx < len(row) else "")
        out.append(item)
    return out


def _read_table_with_fallback(root_dir: str, filename: str) -> List[Dict[str, str]]:
    primary = os.path.join(root_dir, filename)
    rows = _read_table(primary)
    if rows:
        return rows
    parent = os.path.dirname(os.path.abspath(root_dir))
    fallback = os.path.join(parent, filename)
    if os.path.abspath(fallback) == os.path.abspath(primary):
        return rows
    return _read_table(fallback)


def _pick(row: Dict[str, str], keys: List[str], default: str = "") -> str:
    for k in keys:
        if k in row and _norm_text(row.get(k)):
            return _norm_text(row.get(k))
    return default


_TABLE_FILENAMES = (
    "国库.xlsx",
    "内帑.xlsx",
    "地块.xlsx",
    "田税.xlsx",
    "军队.xlsx",
    "海军.xlsx",
    "阶层.xlsx",
    "大臣.xlsx",
    "兵种.xlsx",
    "船种.xlsx",
    "装备.xlsx",
    "国家势力.xlsx",
    "党派.xlsx",
)

_NUMERIC_DATA_CACHE: Dict[str, Tuple[float, NumericData]] = {}


def _normalize_root_dir(root_dir: str) -> str:
    rd = (root_dir or "").strip() or default_data_root()
    if os.path.isabs(rd):
        return os.path.normpath(rd)
    here = Path(__file__).resolve().parent.parent
    return os.path.normpath(str(here / rd))


def _resolve_table_path(root_dir: str, filename: str) -> Optional[str]:
    primary = os.path.join(root_dir, filename)
    if os.path.isfile(primary):
        return os.path.abspath(primary)
    parent = os.path.dirname(os.path.abspath(root_dir))
    fallback = os.path.join(parent, filename)
    if os.path.isfile(fallback):
        return os.path.abspath(fallback)
    return None


def _numeric_data_cache_stamp(root_dir: str) -> float:
    stamp = 0.0
    for filename in _TABLE_FILENAMES:
        path = _resolve_table_path(root_dir, filename)
        if path:
            stamp = max(stamp, os.path.getmtime(path))
    return stamp


def warmup_numeric_data(root_dir: str) -> None:
    """预加载数值表到进程内缓存（开局前调用可隐藏 Excel 读盘耗时）。"""
    load_numeric_data(root_dir)


def load_numeric_data(root_dir: str) -> NumericData:
    abs_root = _normalize_root_dir(root_dir)
    stamp = _numeric_data_cache_stamp(abs_root)
    cached = _NUMERIC_DATA_CACHE.get(abs_root)
    if cached is not None and cached[0] >= stamp:
        return cached[1]

    data = _load_numeric_data(abs_root)
    _NUMERIC_DATA_CACHE[abs_root] = (stamp, data)
    return data


def _load_numeric_data(root_dir: str) -> NumericData:
    data = NumericData()

    treasury_rows = _read_table_with_fallback(root_dir, "国库.xlsx")
    if treasury_rows:
        row = treasury_rows[0]
        data.treasury = TreasurySnapshot(
            balance_liang=_parse_num(row.get("国库余额（两）")),
            total_income_liang=_parse_num(row.get("财政总收入（两）")),
            total_expense_liang=_parse_num(row.get("财政总支出（两）")),
        )

    inner_rows = _read_table_with_fallback(root_dir, "内帑.xlsx")
    if inner_rows:
        row = inner_rows[0]
        data.inner_treasury = TreasurySnapshot(
            balance_liang=_parse_num(row.get("内帑余额（两）")),
            total_income_liang=_parse_num(row.get("内帑总收入（两）")),
            total_expense_liang=_parse_num(row.get("内帑总支出（两）")),
        )

    land_rows = _read_table_with_fallback(root_dir, "地块.xlsx")
    for row in land_rows:
        data.land_blocks.append(
            LandBlock(
                name=row.get("地块名称", ""),
                faction=row.get("所属势力信息", ""),
                total_tax_liang=_parse_num(row.get("总收入税银 (两)")),
                tax_multiplier=_parse_num(row.get("税收倍率"), 1.0) or 1.0,
                tax_completion=_parse_num(row.get("税收完成率"), 1.0) or 1.0,
                public_support=_parse_num(row.get("民心"), 50.0),
                unrest=_parse_num(row.get("动乱"), 0.0),
                disaster_level=row.get("灾害严重等级", ""),
            )
        )

    tax_rows = _read_table_with_fallback(root_dir, "田税.xlsx")
    for row in tax_rows:
        data.tax_rules.append(
            TaxRule(
                land_type=row.get("田亩种类", ""),
                total_tax_per_mu=_parse_num(row.get("田亩总赋税（两 / 亩 / 季度）")),
            )
        )

    army_rows = _read_table_with_fallback(root_dir, "军队.xlsx")
    for row in army_rows:
        data.armies.append(
            ArmyUnit(
                name=row.get("军队名称", ""),
                faction=row.get("所属势力", ""),
                supply=_parse_num(row.get("补给"), 60.0),
                morale=_parse_num(row.get("士气"), 60.0),
                total_people=_parse_num(row.get("总人数")),
                extra_cost_per_person=_parse_num(row.get("额外花费（两 / 人 / 季度）")),
                maintenance_liang=_parse_num(row.get("维护费（两 / 季度）")),
            )
        )

    navy_rows = _read_table_with_fallback(root_dir, "海军.xlsx")
    for row in navy_rows:
        data.navies.append(
            NavyUnit(
                name=row.get("军队名称", ""),
                faction=row.get("所属势力", ""),
                supply=_parse_num(row.get("补给"), 60.0),
                morale=_parse_num(row.get("士气"), 60.0),
                sailors=_parse_num(row.get("水兵人数")),
                maintenance_liang=_parse_num(row.get("维护费（两 / 季度）")),
            )
        )

    strata_rows = _read_table_with_fallback(root_dir, "阶层.xlsx")
    for row in strata_rows:
        data.strata.append(
            StrataGroup(
                name=row.get("阶层名称", ""),
                satisfaction=_parse_num(row.get("满意度"), 50.0),
                influence=_parse_num(row.get("整体影响力"), 50.0),
                revolt_risk=_parse_num(row.get("叛乱风险"), 20.0),
            )
        )

    minister_rows = _read_table_with_fallback(root_dir, "大臣.xlsx")
    for row in minister_rows:
        data.ministers.append(
            Minister(
                name=row.get("姓名", ""),
                office=row.get("官职", ""),
                party=row.get("党派", ""),
                stance=row.get("政治立场", ""),
                influence=_parse_num(row.get("影响力"), 50.0),
                ability_internal=_parse_num(row.get("内政"), 50.0),
                ability_military=_parse_num(row.get("军事"), 50.0),
                ability_diplomacy=_parse_num(row.get("外交"), 50.0),
                ability_finance=_parse_num(row.get("财政"), 50.0),
            )
        )

    army_type_rows = _read_table_with_fallback(root_dir, "兵种.xlsx")
    for row in army_type_rows:
        data.army_types.append(
            ArmyType(
                name=row.get("兵种名称", ""),
                equipment_requirement=row.get("装备需求", ""),
                maintenance_liang=_parse_num(row.get("维护费（两 / 季度）")),
                description=row.get("完整描述", ""),
            )
        )

    ship_type_rows = _read_table_with_fallback(root_dir, "船种.xlsx")
    for row in ship_type_rows:
        data.ship_types.append(
            ShipType(
                name=row.get("船种名称", ""),
                light_cannon_count=_parse_num(row.get("轻型舰炮数量")),
                heavy_cannon_count=_parse_num(row.get("重型舰炮数量")),
                maintenance_liang=_parse_num(row.get("维护费（两 / 季度）")),
                description=row.get("完整描述", ""),
            )
        )

    equipment_rows = _read_table_with_fallback(root_dir, "装备.xlsx")
    for row in equipment_rows:
        data.equipment_items.append(
            EquipmentItem(
                name=row.get("装备名称", ""),
                stock=_parse_num(row.get("库存数量")),
                quarterly_output=_parse_num(row.get("季度产量")),
                description=row.get("描述", ""),
            )
        )

    faction_rows = _read_table_with_fallback(root_dir, "国家势力.xlsx")
    for row in faction_rows:
        data.faction_powers.append(
            FactionPower(
                name=row.get("势力名称", ""),
                ruler=row.get("统治者", ""),
                public_support=_parse_num(row.get("民心"), 50.0),
                population=_parse_num(row.get("总人口")),
                corruption=_parse_num(row.get("贪腐"), 50.0),
                authority=_parse_num(row.get("权威"), 50.0),
                currency_supply=row.get("通货供应", ""),
                controlled_blocks=row.get("控制地块", ""),
            )
        )

    party_rows = _read_table_with_fallback(root_dir, "党派.xlsx")
    for row in party_rows:
        ordered_values = list(row.values())
        name = _pick(row, ["党派名称", "党派", "派系"], default="")
        influence_raw = _pick(row, ["影响力", "整体影响力", "势力影响力"], default="")
        support_raw = _pick(row, ["支持度", "民心", "支持率"], default="")
        stance = _pick(row, ["政治倾向", "立场", "政治立场"], default="")
        description = _pick(row, ["完整描述", "描述", "备注"], default="")

        # Positional fallback for mixed/legacy sheets:
        # col1=name, col2=support, col3=influence, col4=stance, col5=description
        if not name and len(ordered_values) >= 1:
            name = _norm_text(ordered_values[0])
        if not support_raw and len(ordered_values) >= 2:
            support_raw = _norm_text(ordered_values[1])
        if not influence_raw and len(ordered_values) >= 3:
            influence_raw = _norm_text(ordered_values[2])
        if not stance and len(ordered_values) >= 4:
            stance = _norm_text(ordered_values[3])
        if not description and len(ordered_values) >= 5:
            description = _norm_text(ordered_values[4])

        influence = _parse_num(influence_raw, 50.0)
        support = _parse_num(support_raw, 50.0)
        if not name and not row:
            continue
        data.party_groups.append(
            PartyGroup(
                name=name,
                influence=influence,
                support=support,
                stance=stance,
                description=description,
            )
        )

    return data

